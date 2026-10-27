from pathlib import Path
import openpyxl

inp = Path(r"D:/Workspace/iagent/runs/toy/input.xlsx")
out = Path(r"D:/Workspace/iagent/runs/toy/output.xlsx")
wb = openpyxl.load_workbook(inp)
ws = wb.active
ws["C1"] = ws["B1"].value * 2
wb.save(out)
