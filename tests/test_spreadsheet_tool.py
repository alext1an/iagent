from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from iagent.tools.spreadsheet import read_xlsx, summarize_xlsx


def _make_wb(path: Path, *, rows: int = 3, cols: int = 3) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    for r in range(1, rows + 1):
        for c in range(1, cols + 1):
            ws.cell(row=r, column=c, value=f"r{r}c{c}")
    wb.save(path)


def test_summarize_xlsx_reports_sheet_count_names_sizes(tmp_path: Path) -> None:
    p = tmp_path / "multi.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "First"
    ws2 = wb.create_sheet("Second")
    ws2["D4"] = "x"  # force a larger extent
    wb.save(p)
    out = summarize_xlsx(p)
    assert "2 sheet(s): First, Second" in out
    assert "[sheet: 'First']" in out
    assert "[sheet: 'Second']" in out
    assert "max_row=" in out and "max_col=" in out


def test_summarize_xlsx_reads_no_values(tmp_path: Path) -> None:
    p = tmp_path / "secret.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = "TOP_SECRET_VALUE"
    wb.save(p)
    out = summarize_xlsx(p)
    assert "TOP_SECRET_VALUE" not in out


def test_summarize_xlsx_missing_file_raises(tmp_path: Path) -> None:
    with pytest.raises(FileNotFoundError, match="No such workbook"):
        summarize_xlsx(tmp_path / "absent.xlsx")


def test_summarize_xlsx_bad_suffix_raises(tmp_path: Path) -> None:
    p = tmp_path / "data.csv"
    p.write_text("a,b\n", encoding="utf-8")
    with pytest.raises(ValueError, match="Unsupported spreadsheet type"):
        summarize_xlsx(p)


def test_read_xlsx_renders_sheet_and_values(tmp_path: Path) -> None:
    p = tmp_path / "book.xlsx"
    _make_wb(p)
    out = read_xlsx(p)
    assert "workbook: book.xlsx" in out
    assert "[sheet: 'Data']" in out
    assert "r1c1 | r1c2 | r1c3" in out


def test_read_xlsx_blanks_are_empty(tmp_path: Path) -> None:
    p = tmp_path / "gaps.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "x"
    ws["C1"] = "z"  # B1 left blank
    wb.save(p)
    out = read_xlsx(p)
    assert "x |  | z" in out


def test_read_xlsx_truncates_and_notes(tmp_path: Path) -> None:
    p = tmp_path / "big.xlsx"
    _make_wb(p, rows=50, cols=50)
    out = read_xlsx(p, max_rows=5, max_cols=4)
    assert "r5c4" in out
    assert "r6c1" not in out
    assert "rows >5 hidden" in out
    assert "cols >4 hidden" in out


def test_read_xlsx_single_sheet_selection(tmp_path: Path) -> None:
    p = tmp_path / "multi.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "First"
    wb.create_sheet("Second")
    wb.save(p)
    out = read_xlsx(p, sheet="Second")
    assert "[sheet: 'Second']" in out
    assert "[sheet: 'First']" not in out


def test_read_xlsx_missing_sheet_raises(tmp_path: Path) -> None:
    p = tmp_path / "one.xlsx"
    _make_wb(p)
    with pytest.raises(ValueError, match="No sheet named"):
        read_xlsx(p, sheet="Nope")


def test_read_xlsx_missing_file_raises(tmp_path: Path) -> None:
    with pytest.raises(FileNotFoundError, match="No such workbook"):
        read_xlsx(tmp_path / "absent.xlsx")


def test_read_xlsx_bad_suffix_raises(tmp_path: Path) -> None:
    p = tmp_path / "data.csv"
    p.write_text("a,b,c\n", encoding="utf-8")
    with pytest.raises(ValueError, match="Unsupported spreadsheet type"):
        read_xlsx(p)


def test_read_xlsx_nonpositive_bounds_raise(tmp_path: Path) -> None:
    p = tmp_path / "ok.xlsx"
    _make_wb(p)
    with pytest.raises(ValueError, match="must be positive"):
        read_xlsx(p, max_rows=0)
