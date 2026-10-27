from __future__ import annotations

import os
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest
from pydantic_ai.models.test import TestModel
from pydantic_ai_harness import Skills, Shell

from iagent.runtime import load_agent

REPO = Path(__file__).resolve().parents[1]
SPEC = REPO / "src" / "sheetagent" / "agent.yaml"
TOY_DIR = REPO / "runs" / "toy"


def test_sheetagent_yaml_has_expected_fields() -> None:
    text = SPEC.read_text(encoding="utf-8")
    assert "name: sheetagent" in text
    assert "spreadsheet" in text
    assert "LocalWorkspace" in text
    assert "Shell:" in text


@pytest.mark.skipif(
    os.name != "posix",
    reason="LocalWorkspace requires POSIX (use WSL/Linux for full agent load)",
)
def test_sheetagent_yaml_loads_with_skills_and_shell() -> None:
    agent = load_agent(
        SPEC,
        custom_capability_types=[Skills, Shell],
        model=TestModel(custom_output_text="ok", call_tools=[]),
    )
    result = agent.run_sync("hello")
    assert result.output == "ok"


def test_toy_spreadsheet_solution_writes_output() -> None:
    """Run a minimal solution.py like the agent would, without calling the LLM."""
    TOY_DIR.mkdir(parents=True, exist_ok=True)
    input_path = TOY_DIR / "input.xlsx"
    output_path = TOY_DIR / "output.xlsx"
    solution_path = TOY_DIR / "solution.py"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "value"
    ws["B1"] = 10
    wb.save(input_path)

    solution_path.write_text(
        f"""from pathlib import Path
import openpyxl

inp = Path(r"{input_path.as_posix()}")
out = Path(r"{output_path.as_posix()}")
wb = openpyxl.load_workbook(inp)
ws = wb.active
ws["C1"] = ws["B1"].value * 2
wb.save(out)
""",
        encoding="utf-8",
    )

    subprocess.run(
        [sys.executable, str(solution_path)],
        cwd=TOY_DIR,
        check=True,
    )

    assert output_path.is_file()
    out_wb = openpyxl.load_workbook(output_path, read_only=True, data_only=True)
    assert out_wb.active["C1"].value == 20
